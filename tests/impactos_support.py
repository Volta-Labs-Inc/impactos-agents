"""Shared helpers for the impactos CLI test suite (issue #3).

Injects the ``cli/`` package and vendored ``openpyxl`` onto ``sys.path`` and
provides thin wrappers that drive the CLI command functions in-process, plus a
full parse -> apply -> export pipeline over the committed Harbourline fixture.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Dict, List, Optional

REPO_ROOT = Path(__file__).resolve().parents[1]
CLI_DIR = REPO_ROOT / "cli"
VENDOR_DIR = REPO_ROOT / "vendor"
FIXTURE = REPO_ROOT / "fixtures" / "harbourline"
FIXTURE_MAPPINGS = FIXTURE / "mappings"
GOLDEN = FIXTURE / "expected" / "reports" / "2026-06-30"

for path in (str(CLI_DIR), str(VENDOR_DIR)):
    if path not in sys.path:
        sys.path.insert(0, path)

from impactos_agent import cli  # noqa: E402
from impactos_agent import parsers, mapping, workspace  # noqa: E402,F401

PERIOD = "2026-06-30"
DATA_SHEETS = ["Companies", "Company Updates", "Contacts", "Programs", "Program Cohorts"]

CRM_P2 = FIXTURE / "crm" / "period-2026-06-30" / "companies.csv"
CRM_P1 = FIXTURE / "crm" / "period-2026-03-31" / "companies.csv"
PROGRAMS = FIXTURE / "programs" / "harbourline-programs.xlsx"
BAI_WORKBOOK = FIXTURE / "bai-workbook" / "harbourline-bai-v5.xlsx"
CRM_MAPPING = FIXTURE_MAPPINGS / "harbourline-crm.json"
PROGRAMS_MAPPING = FIXTURE_MAPPINGS / "harbourline-programs.json"


def ns(**kwargs) -> SimpleNamespace:
    kwargs.setdefault("json", False)
    return SimpleNamespace(**kwargs)


def run(command: str, **kwargs):
    """Call a CLI command function in-process and return its Result."""
    func = {
        "preflight": cli.cmd_preflight,
        "init": cli.cmd_init,
        "parse": cli.cmd_parse,
        "apply-mapping": cli.cmd_apply_mapping,
        "export": cli.cmd_export,
        "check": cli.cmd_check,
        "state": cli.cmd_state,
        "brief": cli.cmd_brief,
    }[command]
    return func(ns(**kwargs))


def parse_source(project: Path, source: Path, period: str = PERIOD, source_type: Optional[str] = None):
    result = run("parse", source=str(source), type=source_type, period=period, out=None, project_root=project)
    return result


def payload_path(project: Path, source: Path, period: str = PERIOD) -> Path:
    return project / "workspace" / "sources" / period / f"{source.stem}.source.json"


def build_records(project: Path, period: str = PERIOD, include_programs: bool = True) -> Path:
    """Run the full parse -> apply pipeline and return the records.json path."""
    parse_source(project, CRM_P2, period)
    run(
        "apply-mapping",
        source=str(payload_path(project, CRM_P2, period)),
        mapping=str(CRM_MAPPING),
        period=period,
        confirmed=True,
        out=None,
        project_root=project,
    )
    if include_programs:
        parse_source(project, PROGRAMS, period)
        run(
            "apply-mapping",
            source=str(payload_path(project, PROGRAMS, period)),
            mapping=str(PROGRAMS_MAPPING),
            period=period,
            confirmed=True,
            out=None,
            project_root=project,
        )
    return project / "workspace" / "state" / period / "records.json"


def export(project: Path, period: str = PERIOD, out: Optional[Path] = None):
    return run("export", period=period, out=str(out) if out else None, project_root=project)


def report_dir(project: Path, period: str = PERIOD) -> Path:
    return project / "workspace" / "reports" / period


def load_json(path: Path) -> Any:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def workbook_data_cells(path: Path) -> Dict[str, Dict[str, Any]]:
    """Populated cells (row >= 2) on the five data sheets only."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path, data_only=True)
    out: Dict[str, Dict[str, Any]] = {}
    for sheet in DATA_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        cells: Dict[str, Any] = {}
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is not None and v != "":
                    cells[f"{get_column_letter(c)}{r}"] = v
        if cells:
            out[sheet] = cells
    return out
