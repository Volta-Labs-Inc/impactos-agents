"""Deterministic xlsx reading, writing, and BAI-template validation.

Output workbooks are built from the bundled BAI v5 template and canonicalised so
two runs over identical inputs produce byte-identical files: workbook
created/modified timestamps are pinned and every zip entry is rewritten with a
fixed time in sorted order.
"""

from __future__ import annotations

import datetime
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional

from . import paths

paths.inject_vendor()

from openpyxl import load_workbook  # noqa: E402  (import after vendor injection)

# Pinned instants for byte-determinism.
FIXED_DATETIME = datetime.datetime(2026, 1, 1, 0, 0, 0)
FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)
FIXED_ISO = "2026-01-01T00:00:00Z"


def sheet_headers(worksheet) -> List[str]:
    headers = []
    for col in range(1, worksheet.max_column + 1):
        value = worksheet.cell(row=1, column=col).value
        headers.append(str(value).strip() if value is not None else "")
    return headers


def read_tabular_xlsx(path: Path) -> Dict[str, Dict[str, Any]]:
    """Read every sheet of a workbook into ``{sheet: {columns, rows}}``.

    Header cells define the columns; each subsequent non-empty row becomes a dict
    keyed by header. Fully blank rows are skipped.
    """
    workbook = load_workbook(path, data_only=True)
    sheets: Dict[str, Dict[str, Any]] = {}
    try:
        for worksheet in workbook.worksheets:
            headers = [h for h in sheet_headers(worksheet) if h]
            rows: List[Dict[str, Any]] = []
            for excel_row in worksheet.iter_rows(min_row=2, values_only=True):
                values = list(excel_row)
                if all(value is None or value == "" for value in values):
                    continue
                row: Dict[str, Any] = {}
                for index, header in enumerate(headers):
                    row[header] = values[index] if index < len(values) else None
                rows.append(row)
            sheets[worksheet.title] = {"columns": headers, "rows": rows}
    finally:
        workbook.close()
    return sheets


def _pin_core_date(data: bytes, tag: bytes) -> bytes:
    """Replace only the text inside ``<tag ...>TEXT</tag>`` with the fixed instant."""
    pattern = rb"(<" + re.escape(tag) + rb"\b[^>]*>).*?(</" + re.escape(tag) + rb">)"
    return re.sub(pattern, rb"\g<1>" + FIXED_ISO.encode() + rb"\g<2>", data)


def canonicalize_xlsx(path: Path) -> None:
    """Rewrite an xlsx in place with pinned timestamps and sorted entries."""
    temporary = path.with_suffix(path.suffix + ".canon")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            data = source.read(name)
            if name == "docProps/core.xml":
                # openpyxl resets <dcterms:modified> to the wall clock on save.
                # Pin it to the fixed instant, preserving the element's inline
                # namespace attributes (only the text between the tags changes).
                data = _pin_core_date(data, b"dcterms:modified")
                data = _pin_core_date(data, b"dcterms:created")
            target.writestr(info, data)
    temporary.replace(path)


def safe_excel_value(value: Any) -> Any:
    """Neutralise a leading formula character so a cell is never a live formula."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def build_workbook(sheet_rows: Dict[str, List[List[Any]]], destination: Path) -> None:
    """Build an output workbook from the bundled template and canonicalise it.

    ``sheet_rows`` maps a template sheet name to the data rows (excluding the
    header) to write beneath its existing header. Sheets not listed keep the
    template's header only. Template sheets, headers and dropdowns are preserved.
    """
    workbook = load_workbook(paths.BUNDLED_TEMPLATE)
    for sheet_name, rows in sheet_rows.items():
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"template has no sheet {sheet_name!r}")
        worksheet = workbook[sheet_name]
        for row_offset, row in enumerate(rows, start=2):
            for col_offset, value in enumerate(row, start=1):
                worksheet.cell(row=row_offset, column=col_offset, value=safe_excel_value(value))
    workbook.properties.created = FIXED_DATETIME
    workbook.properties.modified = FIXED_DATETIME
    workbook.properties.creator = "impactos"
    workbook.properties.lastModifiedBy = "impactos"
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    canonicalize_xlsx(destination)


def validate_bai_template(path: Path, export_contract: Dict[str, Any]) -> List[str]:
    """Check a workbook against the BAI v5 export contract.

    Ported from the BAI template validator: every required sheet is present with
    the exact column headers in order, and each declared Excel dropdown exists on
    its range with the exact vocabulary. Returns a list of problems (empty = ok).
    """
    problems: List[str] = []
    workbook = load_workbook(path, data_only=False)
    sheets = export_contract["sheets"]
    for sheet_name, sheet_spec in sheets.items():
        if sheet_name not in workbook.sheetnames:
            problems.append(f"missing sheet: {sheet_name}")
            continue
        worksheet = workbook[sheet_name]
        expected = [column["column"] for column in sheet_spec["columns"]]
        actual = [worksheet.cell(1, index + 1).value for index in range(len(expected))]
        if actual != expected:
            problems.append(f"header mismatch on {sheet_name}: {actual} != {expected}")

    vocabularies = export_contract["vocabularies"]
    for validation in export_contract.get("excel_validations", []):
        sheet_name = validation["sheet"]
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        expected_values = vocabularies[validation["vocabulary"]]
        expected_formula = '"%s"' % ",".join(expected_values)
        matches = [
            item
            for item in worksheet.data_validations.dataValidation
            if item.type == "list"
            and str(item.sqref) == validation["range"]
            and item.formula1 == expected_formula
        ]
        if not matches:
            problems.append(
                f"missing/incorrect dropdown on {sheet_name}.{validation['field']} ({validation['range']})"
            )
    return problems


def read_all_cells(path: Path) -> Dict[str, Dict[str, Any]]:
    """Return every populated data cell (row >= 2) as ``{sheet: {A2: value}}``."""
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path, data_only=True)
    result: Dict[str, Dict[str, Any]] = {}
    for worksheet in workbook.worksheets:
        cells: Dict[str, Any] = {}
        for row in range(2, worksheet.max_row + 1):
            for col in range(1, worksheet.max_column + 1):
                value = worksheet.cell(row=row, column=col).value
                if value is not None and value != "":
                    cells[f"{get_column_letter(col)}{row}"] = value
        result[worksheet.title] = cells
    return result
